"""local 模式 Office 文档（docx/pptx/xlsx）纯 Python 文本抽取。

只做「能被检索的正文」级别的抽取：段落/表格 → Markdown，不做版式还原。
旧版二进制格式（.doc/.ppt/.xls）不在支持范围（无纯 Python 可靠解析库）。
"""

from __future__ import annotations

import io

# 防止超大表格撑爆文档正文：每表最多抽取的行/列数
_MAX_SHEET_ROWS = 500
_MAX_SHEET_COLS = 50


def _is_heading_style(style_name: str) -> int | None:
    """python-docx 标题样式名（如 Heading 1 / 标题 1）→ Markdown 层级；非标题返回 None。"""
    lowered = (style_name or "").lower()
    for token in ("heading", "标题"):
        idx = lowered.find(token)
        if idx >= 0:
            digits = "".join(ch for ch in lowered[idx + len(token):] if ch.isdigit())
            if digits:
                return min(int(digits), 6)
    return None


def extract_docx(raw: bytes) -> str:
    """docx → Markdown：段落（标题加 #、列表加 -）+ 表格转 Markdown 表。"""
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(raw))
    except Exception as exc:  # noqa: BLE001 — 损坏/加密文件转用户可读错误
        raise ValueError(f"docx 解析失败（文件可能损坏或已加密）：{exc}") from exc

    # 按文档体真实顺序遍历段落与表格，保持阅读顺序
    body = doc.element.body
    parts: list[str] = []
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            level = _is_heading_style(para.style.name if para.style is not None else "")
            if level:
                parts.append(f"{'#' * level} {text}")
            elif para.style is not None and "list" in (para.style.name or "").lower():
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif isinstance(child, CT_Tbl):
            # XML 元素需包成 Table 代理才能按 row/cell 读取
            table = Table(child, doc)
            parts.append(_md_table([[c.text.strip() for c in row.cells] for row in table.rows]))
    return "\n\n".join(parts).strip()


def _md_table(rows: list[list[str]]) -> str:
    """二维数组 → Markdown 表；空表返回空串。"""
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    normalized = [r + [""] * (width - len(r)) for r in rows]
    head = "| " + " | ".join(normalized[0]) + " |"
    sep = "| " + " | ".join(["---"] * width) + " |"
    body = "\n".join("| " + " | ".join(r) + " |" for r in normalized[1:])
    return "\n".join([head, sep, body]) if body else "\n".join([head, sep])


def extract_pptx(raw: bytes) -> str:
    """pptx → 逐页「## Slide N」+ 文本框/表格/演讲者备注。"""
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(raw))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"pptx 解析失败（文件可能损坏或已加密）：{exc}") from exc

    parts: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        page: list[str] = [f"## Slide {idx}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
                if texts:
                    page.extend(texts)
            if getattr(shape, "has_table", False):
                rows = [[c.text.strip() for c in row.cells] for row in shape.table.rows]
                table_md = _md_table(rows)
                if table_md:
                    page.append(table_md)
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                page.append(f"> 备注：{notes}")
        # 跳过整页无内容的空白页（纯图片页等）
        if len(page) > 1:
            parts.append("\n".join(page))
    return "\n\n".join(parts).strip()


def extract_xlsx(raw: bytes) -> str:
    """xlsx → 逐工作表「## 工作表：名」+ TSV 行（限量防超大文件撑爆正文）。"""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"xlsx 解析失败（文件可能损坏或已加密）：{exc}") from exc

    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"## 工作表：{ws.title}")
        rows_out = 0
        for row in ws.iter_rows(max_row=_MAX_SHEET_ROWS, max_col=_MAX_SHEET_COLS, values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if not any(c.strip() for c in cells):
                continue
            parts.append("\t".join(cells))
            rows_out += 1
            if rows_out >= _MAX_SHEET_ROWS:
                parts.append(f"…（工作表超过 {_MAX_SHEET_ROWS} 行，已截断）")
                break
    return "\n".join(parts).strip()


def extract_office_text(ext: str, raw: bytes) -> str:
    """按扩展名分发 Office 文本抽取；不支持的扩展名抛 ValueError。"""
    extractors = {
        ".docx": extract_docx,
        ".pptx": extract_pptx,
        ".xlsx": extract_xlsx,
    }
    extractor = extractors.get(ext)
    if extractor is None:
        raise ValueError(f"不支持的 Office 格式：{ext}")
    return extractor(raw)
