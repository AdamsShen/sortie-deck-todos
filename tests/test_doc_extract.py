"""local 模式 Office 文档解析（docx/pptx/xlsx）测试：内存生成真实文档 → 抽取断言。"""

from __future__ import annotations

import io

import pytest

from sortie_deck.doc_extract import extract_docx, extract_office_text, extract_pptx, extract_xlsx


def _make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("作战手册", level=1)
    doc.add_paragraph("第一段正文内容 marker-docx-p1")
    doc.add_paragraph("列表项甲", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列A"
    table.cell(0, 1).text = "列B"
    table.cell(1, 0).text = "v1"
    table.cell(1, 1).text = "v2"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "季度汇报 marker-pptx-title"
    slide.placeholders[1].text = "要点一"
    slide_layout = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide_layout.shapes.add_textbox(0, 0, 100, 50)
    tb.text_frame.text = "第二页文本 marker-pptx-p2"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "对账"
    ws.append(["日期", "金额"])
    ws.append(["2026-08-01", "1200.5"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_docx_paragraph_heading_table():
    text = extract_docx(_make_docx())
    assert "# 作战手册" in text
    assert "marker-docx-p1" in text
    assert "- 列表项甲" in text
    assert "| 列A | 列B |" in text and "| v1 | v2 |" in text


def test_extract_pptx_slides_and_textboxes():
    text = extract_pptx(_make_pptx())
    assert "## Slide 1" in text
    assert "marker-pptx-title" in text
    assert "marker-pptx-p2" in text


def test_extract_xlsx_sheets_and_rows():
    text = extract_xlsx(_make_xlsx())
    assert "## 工作表：对账" in text
    assert "2026-08-01\t1200.5" in text


def test_dispatch_by_ext_and_unsupported():
    assert "marker-docx-p1" in extract_office_text(".docx", _make_docx())
    assert "marker-pptx-p2" in extract_office_text(".pptx", _make_pptx())
    assert "2026-08-01" in extract_office_text(".xlsx", _make_xlsx())
    with pytest.raises(ValueError):
        extract_office_text(".doc", b"legacy binary")


def test_corrupted_files_raise_value_error():
    for ext in (".docx", ".pptx", ".xlsx"):
        with pytest.raises(ValueError):
            extract_office_text(ext, b"not a real office file")
